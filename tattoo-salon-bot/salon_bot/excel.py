"""Excel exports of the shift schedule and the month's expenses."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import DayOff, Expense, Master, Shift
from .fmt import WEEKDAYS_SHORT, category_title, days_between, month_bounds, month_title
from .payments import Occurrence

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="EDEDEA")
SHIFT_FILL = PatternFill("solid", fgColor="CDE2FB")
CLOSED_FILL = PatternFill("solid", fgColor="F0EFEC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONEY = '#,##0.00 "₽"'


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def schedule_xlsx(
    year: int,
    month: int,
    masters: list[Master],
    shifts: list[Shift],
    days_off: list[DayOff],
    open_weekdays: set[int],
) -> bytes:
    days = days_between(*month_bounds(year, month))
    by_key = {(s.master_id, s.day): s for s in shifts}
    off = {(o.master_id, o.day) for o in days_off}
    known = {m.id for m in masters}
    rows = [(m.id, m.name) for m in masters] + sorted(
        {s.master_id: s.master_name for s in shifts if s.master_id not in known}.items()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "График"
    ws.cell(1, 1, f"График работы — {month_title(year, month)}").font = Font(bold=True, size=13)
    ws.cell(2, 1, "Мастер").font = BOLD
    for c, day in enumerate(days, start=2):
        cell = ws.cell(2, c, f"{day.day}\n{WEEKDAYS_SHORT[day.weekday()]}")
        cell.font, cell.alignment, cell.fill = BOLD, CENTER, HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = 6.5
    total_col = len(days) + 2
    ws.cell(2, total_col, "Смен").font = BOLD
    ws.cell(2, total_col + 1, "Часов").font = BOLD
    ws.row_dimensions[2].height = 30
    ws.column_dimensions["A"].width = 18

    for r, (master_id, name) in enumerate(rows, start=3):
        ws.cell(r, 1, name)
        count, minutes = 0, 0
        for c, day in enumerate(days, start=2):
            cell = ws.cell(r, c)
            cell.alignment = CENTER
            if shift := by_key.get((master_id, day)):
                cell.value = f"{shift.start}\n{shift.end}"
                cell.fill = SHIFT_FILL
                count += 1
                minutes += _minutes(shift.end) - _minutes(shift.start)
            elif (master_id, day) in off:
                cell.value = "вых"
            elif day.weekday() not in open_weekdays:
                cell.fill = CLOSED_FILL
        ws.cell(r, total_col, count)
        ws.cell(r, total_col + 1, round(minutes / 60, 1))
        ws.row_dimensions[r].height = 30
    ws.freeze_panes = "B3"

    flat = wb.create_sheet("Списком")
    flat.append(["Дата", "День", "Мастер", "Начало", "Конец"])
    for cell in flat[1]:
        cell.font = BOLD
    for s in sorted(shifts, key=lambda s: (s.day, s.start, s.master_name)):
        flat.append([s.day, WEEKDAYS_SHORT[s.day.weekday()], s.master_name, s.start, s.end])
        flat.cell(flat.max_row, 1).number_format = "DD.MM.YYYY"
    for col, width in zip("ABCDE", (12, 6, 18, 8, 8)):
        flat.column_dimensions[col].width = width
    return _to_bytes(wb)


def expenses_xlsx(
    year: int, month: int, expenses: list[Expense], plan: list[Occurrence], today: date
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Расходы"
    ws.append(["Дата", "Категория", "Сумма", "Комментарий"])
    for e in expenses:
        ws.append([e.day, category_title(e.category), e.amount / 100, e.comment])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 3).number_format = MONEY
    total_row = ws.max_row + 1
    ws.cell(total_row, 2, "Итого").font = BOLD
    total = ws.cell(total_row, 3, f"=SUM(C2:C{max(total_row - 1, 2)})")
    total.font, total.number_format = BOLD, MONEY
    for col, width in zip("ABCD", (12, 22, 14, 40)):
        ws.column_dimensions[col].width = width

    cats = wb.create_sheet("По категориям")
    cats.append(["Категория", "Сумма", "Доля"])
    totals: dict[str, int] = {}
    for e in expenses:
        totals[e.category] = totals.get(e.category, 0) + e.amount
    grand = sum(totals.values())
    for code, amount in sorted(totals.items(), key=lambda kv: -kv[1]):
        cats.append([category_title(code), amount / 100, amount / grand if grand else 0])
        cats.cell(cats.max_row, 2).number_format = MONEY
        cats.cell(cats.max_row, 3).number_format = "0%"
    cats.column_dimensions["A"].width = 22
    cats.column_dimensions["B"].width = 14

    pay = wb.create_sheet("График платежей")
    pay.append(["Срок", "Платёж", "Категория", "Сумма", "Статус"])
    for o in plan:
        status = "оплачено" if o.paid else ("просрочено" if o.due < today else "предстоит")
        pay.append([o.due, o.planned.title, category_title(o.planned.category), o.planned.amount / 100, status])
        pay.cell(pay.max_row, 1).number_format = "DD.MM.YYYY"
        pay.cell(pay.max_row, 4).number_format = MONEY
    for col, width in zip("ABCDE", (12, 26, 22, 14, 12)):
        pay.column_dimensions[col].width = width

    for sheet in (ws, cats, pay):
        for cell in sheet[1]:
            cell.font, cell.fill = BOLD, HEADER_FILL
        sheet.freeze_panes = "A2"
    ws.title = f"Расходы {month:02d}.{year}"
    return _to_bytes(wb)


def _minutes(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)
